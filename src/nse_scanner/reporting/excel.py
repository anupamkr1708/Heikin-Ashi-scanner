"""Excel report generation (PART 44), with a deliberate fix for BUG 14 / BUG 7.

**Audit note (BUG 14):** the legacy notebook could turn an already-scaled percentage-point value
(e.g. `BB_Overshoot_Pct = 2.31`, meaning "2.31%") into "231%" in Excel, because Excel's native
`0.00%` number format ALSO multiplies the underlying number by 100 for display — applying it to
a value that was already multiplied by 100 double-counts.

**Resolution used here (documented, not silent):** every `*_Pct` field in this codebase (e.g.
`BB_Overshoot_Pct`, `HA_Body_Pct`, `Dist_SMA50_Pct`) is a percentage-POINT value by construction
— the mandatory baseline condition is literally written as `0 < BB_Overshoot_Pct <= 4.0`
(PART 1), which only makes sense if 4.0 means "4 percent", not "400 percent". Changing that
internal convention to a 0-1 fraction (as PART 44's example literally suggests) would silently
change the meaning of the immutable baseline's own threshold constants — worse than the bug it's
fixing. Instead, `*_Pct` columns get a CUSTOM Excel number format, `0.00"%"`, which appends a
literal percent sign for display WITHOUT re-multiplying by 100. `BB_PctB` is a RATIO (PART 17),
not a percentage, and is formatted as a plain 4-decimal number — never with a `%` sign.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PCT_POINT_FORMAT = '0.00"%"'      # for *_Pct fields already scaled to percentage points
RATIO_FORMAT = "0.0000"           # for BB_PctB and similar true ratios
PRICE_FORMAT = "0.00"
INT_FORMAT = "0"

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)

RATIO_COLUMN_NAMES = {"BB_PctB"}


def _infer_number_format(column_name: str) -> str | None:
    if column_name in RATIO_COLUMN_NAMES:
        return RATIO_FORMAT
    if column_name.endswith("_Pct"):
        return PCT_POINT_FORMAT
    if column_name in ("CMP", "Close", "Open", "High", "Low", "Entry_Price", "Signal_Close",
                        "BB_Upper", "BB_Middle", "BB_Lower", "ATR14"):
        return PRICE_FORMAT
    if column_name in ("Volume", "Rank", "Days_Above_Upper_BB", "N"):
        return INT_FORMAT
    return None


def _write_sheet(wb: openpyxl.Workbook, sheet_name: str, df: pd.DataFrame) -> None:
    ws = wb.create_sheet(title=sheet_name[:31])
    if df is None or df.empty:
        ws["A1"] = f"(no data for {sheet_name})"
        return

    for col_idx, col_name in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=str(col_name))
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")

    formats = {col: _infer_number_format(col) for col in df.columns}

    for row_idx, row in enumerate(df.itertuples(index=False), start=2):
        for col_idx, (col_name, value) in enumerate(zip(df.columns, row, strict=False), start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=_safe_value(value))
            fmt = formats.get(col_name)
            if fmt:
                cell.number_format = fmt

    for col_idx, col_name in enumerate(df.columns, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = max(12, min(28, len(str(col_name)) + 4))
    ws.freeze_panes = "A2"


def _safe_value(value):
    if pd.isna(value):
        return None
    if isinstance(value, (pd.Timestamp,)):
        return value.to_pydatetime()
    return value


@dataclass
class ReportSheets:
    live_signals: pd.DataFrame = field(default_factory=pd.DataFrame)
    stale_signals: pd.DataFrame = field(default_factory=pd.DataFrame)
    diagnostics: pd.DataFrame = field(default_factory=pd.DataFrame)
    data_health: pd.DataFrame = field(default_factory=pd.DataFrame)
    scan_log: pd.DataFrame = field(default_factory=pd.DataFrame)
    universe: pd.DataFrame = field(default_factory=pd.DataFrame)
    parameters: pd.DataFrame = field(default_factory=pd.DataFrame)
    market_regime: pd.DataFrame = field(default_factory=pd.DataFrame)
    research_summary: pd.DataFrame = field(default_factory=pd.DataFrame)
    forward_returns: pd.DataFrame = field(default_factory=pd.DataFrame)
    sensitivity: pd.DataFrame = field(default_factory=pd.DataFrame)
    readme_text: str = ""


def write_report(sheets: ReportSheets, output_path: str | Path) -> Path:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # drop the default empty sheet

    _write_sheet(wb, "Live_Signals", sheets.live_signals)
    _write_sheet(wb, "Stale_Signals", sheets.stale_signals)
    _write_sheet(wb, "Diagnostics", sheets.diagnostics)
    _write_sheet(wb, "Data_Health", sheets.data_health)
    _write_sheet(wb, "Scan_Log", sheets.scan_log)
    _write_sheet(wb, "Universe", sheets.universe)
    _write_sheet(wb, "Parameters", sheets.parameters)
    _write_sheet(wb, "Market_Regime", sheets.market_regime)
    _write_sheet(wb, "Research_Summary", sheets.research_summary)
    _write_sheet(wb, "Forward_Returns", sheets.forward_returns)
    _write_sheet(wb, "Sensitivity", sheets.sensitivity)

    ws = wb.create_sheet(title="README")
    ws["A1"] = "NSE EOD/T-1 Technical Scanner — Report Notes"
    ws["A1"].font = Font(bold=True, size=14)
    for i, line in enumerate((sheets.readme_text or _default_readme()).split("\n"), start=3):
        ws.cell(row=i, column=1, value=line)
    ws.column_dimensions["A"].width = 110

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return Path(output_path)


def _default_readme() -> str:
    return (
        "This is an NSE EOD/T-1 Technical Scanner report — NOT a live/intraday signal feed.\n"
        "See Data_Health for whether this run is trustworthy (RUN_HEALTH = GREEN/YELLOW/RED).\n"
        "Only rows with Data_Status = CURRENT appear in Live_Signals; everything else is in "
        "Stale_Signals.\n"
        "*_Pct columns are percentage-POINT values (e.g. 2.31 means 2.31%), formatted with a "
        "literal '%' suffix — they are not Excel-native percentages and are not re-scaled.\n"
        "BB_PctB is a RATIO (0=lower band, 1=upper band), not a percentage.\n"
        "Research_Heuristic_Score (if present) is NOT a probability and is NOT statistically "
        "validated — see Research_Summary and RESEARCH_METHODOLOGY.md in the repository.\n"
    )
