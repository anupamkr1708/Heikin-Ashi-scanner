import openpyxl
import pandas as pd
from nse_scanner.reporting.excel import (
    PCT_POINT_FORMAT,
    RATIO_FORMAT,
    ReportSheets,
    _infer_number_format,
    write_report,
)


def test_pct_columns_get_literal_percent_format_not_native_percent():
    # BUG 14: a native Excel '0.00%' format would multiply 2.31 by 100 and show 231%.
    fmt = _infer_number_format("BB_Overshoot_Pct")
    assert fmt == PCT_POINT_FORMAT
    assert "%" not in fmt.replace('"%"', "")  # the % is a literal suffix, not the native token


def test_pctb_gets_ratio_format_not_percent():
    fmt = _infer_number_format("BB_PctB")
    assert fmt == RATIO_FORMAT
    assert "%" not in fmt


def test_write_report_roundtrip_preserves_value_magnitude(tmp_path):
    df = pd.DataFrame([{"Stock": "FOO", "BB_Overshoot_Pct": 2.31, "BB_PctB": 1.0575}])
    sheets = ReportSheets(live_signals=df)
    out_path = tmp_path / "report.xlsx"
    write_report(sheets, out_path)

    wb = openpyxl.load_workbook(out_path)
    ws = wb["Live_Signals"]
    # Header row 1, data row 2. Find BB_Overshoot_Pct column.
    headers = [c.value for c in ws[1]]
    col = headers.index("BB_Overshoot_Pct") + 1
    cell = ws.cell(row=2, column=col)
    # The underlying stored number must remain 2.31 (NOT 0.0231) — never re-scaled.
    assert cell.value == 2.31
    assert cell.number_format == PCT_POINT_FORMAT
